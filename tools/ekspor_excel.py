# -*- coding: utf-8 -*-
"""
Ekspor titik reklame ke berkas Excel dengan kolom terpilih.

    python tools/ekspor_excel.py [keluaran.xlsx]

Koordinat diambil dari geometri, bukan dari kolom LONG/LAT, karena
geometri itulah yang dipakai peta dan yang menjadi dasar hitungan jarak POV.
"""
import json
import os
import sys
import math
from datetime import datetime

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")

FONT = "Arial"

# (judul kolom, nama kolom di data, format angka, lebar)
KOLOM = [
    ("UID",                  "UID",         None,          11),
    ("ID Titik",             "ID_TITIK",    None,          24),
    ("Kewenangan",           "KEWENANGAN",  None,          17),
    ("Lintang (LAT)",        "_LAT",        "0.00000000",  15),
    ("Bujur (LONG)",         "_LON",        "0.00000000",  15),
    ("Nama Jalan",           "NAMA_JALAN",  None,          38),
    ("Kelas Jalan",          "KELAS_JLN",   None,          12),
    ("ROW Jalan (m)",        "ROW_JALAN",   "0.0",         14),
    ("POV Rata-rata (m)",    "POV_RERATA",  "0.0",         17),
    ("POV Median (m)",       "POV_MEDIAN",  "0.0",         15),
    ("POV Maksimum (m)",     "POV_MAKS",    "0.0",         17),
    ("POV Minimum (m)",      "POV_MIN",     "0.0",         16),
    ("Kelurahan",            "KELURAHAN",   None,          20),
    ("Kecamatan",            "KECAMATAN",   None,          16),
    ("Sub Wilayah SK 321",   "STATUS_WPP",  None,          26),
    ("Jenis",                "JENIS",       None,          14),
    ("Tipe",                 "TIPE",        None,          22),
    ("Tipe Media",           "TIPE_MEDIA",  None,          14),
]


def jarak(lon1, lat1, lon2, lat2):
    R = 6371008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp, dl = p2 - p1, math.radians(lon2 - lon1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * R * math.asin(math.sqrt(h))


def main():
    keluar = sys.argv[1] if len(sys.argv) > 1 else os.path.join(
        ROOT, "keluaran", "Titik_Reklame_Batam.xlsx")
    os.makedirs(os.path.dirname(keluar), exist_ok=True)

    gj = json.load(open(os.path.join(DATA, "reklame.geojson"), encoding="utf-8"))
    feats = gj["features"]

    baris = []
    beda_koord = 0
    for f in feats:
        p = dict(f["properties"])
        g = f["geometry"]["coordinates"] if f["geometry"] else [None, None]
        p["_LON"], p["_LAT"] = g[0], g[1]
        if g[0] is not None and p.get("LONG") is not None and p.get("LAT") is not None:
            if jarak(g[0], g[1], p["LONG"], p["LAT"]) > 1:
                beda_koord += 1
        baris.append(p)

    wb = Workbook()
    ws = wb.active
    ws.title = "Titik Reklame"

    kepala_isi = PatternFill("solid", fgColor="1F4E79")
    kepala_font = Font(name=FONT, size=10, bold=True, color="FFFFFF")
    garis = Side(style="thin", color="BFBFBF")
    tepi = Border(left=garis, right=garis, top=garis, bottom=garis)

    for i, (judul, _kunci, _fmt, lebar) in enumerate(KOLOM, start=1):
        c = ws.cell(row=1, column=i, value=judul)
        c.font = kepala_font
        c.fill = kepala_isi
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = tepi
        ws.column_dimensions[get_column_letter(i)].width = lebar
    ws.row_dimensions[1].height = 30

    isi_font = Font(name=FONT, size=10)
    for r, p in enumerate(baris, start=2):
        for i, (_judul, kunci, fmt, _lebar) in enumerate(KOLOM, start=1):
            c = ws.cell(row=r, column=i, value=p.get(kunci))
            c.font = isi_font
            c.border = tepi
            if fmt:
                c.number_format = fmt
                c.alignment = Alignment(horizontal="right")

    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(KOLOM)), len(baris) + 1)

    # ---------------------------------------------------------- keterangan
    k = wb.create_sheet("Keterangan")
    k.column_dimensions["A"].width = 30
    k.column_dimensions["B"].width = 96
    judul = Font(name=FONT, size=12, bold=True)
    tebal = Font(name=FONT, size=10, bold=True)
    biasa = Font(name=FONT, size=10)

    terisi = lambda kunci: sum(1 for p in baris if p.get(kunci) is not None)
    catatan = [
        ("Berkas", "Titik reklame Kota Batam"),
        ("Jumlah titik", len(baris)),
        ("Dibuat", datetime.now().strftime("%d %B %Y, %H:%M")),
        ("Sumber", "data/reklame.geojson pada repositori WebGIS reklame4"),
        ("", ""),
        ("Koordinat", "Diambil dari geometri titik, bukan dari kolom LONG/LAT pada tabel "
                      "atribut. Geometri inilah yang dipakai peta dan menjadi dasar seluruh "
                      "hitungan jarak POV. Sistem koordinat WGS 84 (EPSG:4326), derajat desimal."),
        ("Selisih koordinat", "%d titik yang kolom LONG/LAT-nya masih berbeda lebih dari 1 m "
                              "dari geometrinya. Kolom pada berkas ini memakai nilai geometri."
                              % beda_koord),
        ("", ""),
        ("Kolom POV", "POV Rata-rata, Median, Maksimum, dan Minimum adalah jarak dari titik "
                      "reklame ke titik-titik POV miliknya, dalam meter. Dihitung ulang dari "
                      "geometri, bukan dari nilai lama yang tersimpan."),
        ("POV kosong", "Sel POV yang kosong berarti titik tersebut belum memiliki titik POV. "
                       "%d dari %d titik memiliki POV." % (terisi("POV_RERATA"), len(baris))),
        ("", ""),
        ("ROW Jalan", "Lebar ruang milik jalan hasil pengukuran tim GIS, dalam meter. "
                      "Terisi pada %d titik." % terisi("ROW_JALAN")),
        ("Sub Wilayah SK 321", "Berasal dari kolom STATUS_WPP. Terisi pada %d titik."
                               % terisi("STATUS_WPP")),
        ("", ""),
        ("Catatan", "Baris diurutkan sama seperti pada berkas GIS agar mudah digabungkan "
                    "kembali. Gunakan penyaring pada baris judul untuk mengurutkan sendiri. "
                    "UID adalah penanda yang tidak pernah berubah; ID Titik dapat berganti "
                    "bila terjadi penomoran ulang."),
    ]
    k.cell(row=1, column=1, value="Keterangan Berkas").font = judul
    for i, (a, b) in enumerate(catatan, start=3):
        ca = k.cell(row=i, column=1, value=a)
        ca.font = tebal
        ca.alignment = Alignment(vertical="top")
        cb = k.cell(row=i, column=2, value=b)
        cb.font = biasa
        cb.alignment = Alignment(vertical="top", wrap_text=True)

    wb.save(keluar)
    print("tersimpan: %s" % keluar)
    print("  %d baris x %d kolom, %.2f MB"
          % (len(baris), len(KOLOM), os.path.getsize(keluar) / 1e6))
    print("  titik ber-POV: %d | ROW terisi: %d | koordinat berbeda >1 m: %d"
          % (terisi("POV_RERATA"), terisi("ROW_JALAN"), beda_koord))


if __name__ == "__main__":
    main()
