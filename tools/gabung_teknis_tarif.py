# -*- coding: utf-8 -*-
"""
Bangun ulang data reklame dari dua sumber resmi:

  A. GeoPackage "Technical Only" dari tim GIS  -> geometri + kolom teknis
  B. Kertas Kerja Tarif (Excel)                -> seluruh kolom hitungan

    python tools/gabung_teknis_tarif.py <teknis.gpkg> <tarif.xlsx>

UID menjadi kunci penggabungan. Nama kolom Excel diterjemahkan ke kode
kolom lewat Kamus_Kolom, sisanya memakai peta KODE_TAMBAHAN di bawah.

Kolom ringkasan POV pada Excel adalah potret per 2 September, sebelum
127 titik POV baru ditambahkan. Nilai itu tetap dibawa sebagai dasar
perhitungan tarif dengan akhiran _T8, sedangkan kolom POV utama diisi
hasil hitungan terkini dari lapisan POV.
"""
import json
import os
import sys
import sqlite3
from collections import OrderedDict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
sys.path.insert(0, HERE)
from gpkg_to_geojson import gpkg_point, baca_kamus, bangun_skema, GRUP, GRUP_LAIN, SRC_KAMUS

KUNCI_XL = "UID GIS"

# Kolom Excel yang tidak tercantum di Kamus_Kolom -> kode kolom di data spasial.
KODE_TAMBAHAN = {
    "Jumlah sisi": ("JML_MUKA", "Jumlah sisi"),
    "Luas lahan yang diduduki m2": ("LUAS_NAUNG", "Luas lahan yang diduduki m2"),
    "Batas bawah nilai tanah Rp/th": ("LANTAI_TN", "Batas bawah nilai tanah Rp/th"),
    "Kisaran CPM bawah": ("PITA_CPMB", "Kisaran CPM bawah"),
    "Kisaran CPM atas": ("PITA_CPMA", "Kisaran CPM atas"),
    "Kisaran ROW (dari ROW terukur)": ("PITA_ROW", "Kisaran ROW"),
    "Ukuran koridor menurut kisaran ROW": ("UKR_KORID", "Ukuran koridor menurut kisaran ROW"),
    "Jumlah ruas ZNT yang dibandingkan": ("RUAS_ZNT", "Jumlah ruas ZNT dibandingkan"),
}
# Lima keluarga kolom untuk tujuh susunan NS.
for _ns in ["1", "A", "B", "4", "5", "C", "D"]:
    KODE_TAMBAHAN["SEWA NS " + _ns] = ("SEWA_NS" + _ns, "Sewa NS %s (Rp/th)" % _ns)
    KODE_TAMBAHAN["CPM NS " + _ns] = ("CPM_NS" + _ns, "CPM wajib NS %s" % _ns)
    KODE_TAMBAHAN["Status kisaran NS " + _ns] = ("KDPITA_NS" + _ns,
                                                 "Status kisaran NS %s" % _ns)
    KODE_TAMBAHAN["Status NS " + _ns] = ("STAT_NS" + _ns, "Status NS %s" % _ns)
    KODE_TAMBAHAN["Batas atas NS " + _ns] = ("PAGAR_NS" + _ns, "Batas atas NS %s" % _ns)

# Kolom Excel yang sudah diwakili sumber teknis atau geometri: tidak dibawa.
ABAIKAN = {"UID GIS", "ID Titik", "Kecamatan", "Kelurahan", "Nama jalan",
           "Kewenangan", "Sub wilayah SK 321", "Tipe kawasan",
           "ROW terukur tim GIS (m)", "Bujur", "Lintang"}

# Potret POV yang dipakai kertas kerja tarif; kolom POV utama tetap dari
# hitungan terkini agar sesuai dengan lapisan POV.
POV_POTRET = {
    "POV rerata GIS 02-09 (m)": ("POV_RER_T8", "POV rerata dasar Tarif 8 (m)"),
    "POV maksimum GIS 02-09 (m)": ("POV_MAK_T8", "POV maksimum dasar Tarif 8 (m)"),
    "POV minimum GIS 02-09 (m)": ("POV_MIN_T8", "POV minimum dasar Tarif 8 (m)"),
    "Jumlah titik POV GIS": ("JML_POV_T8", "Jumlah POV dasar Tarif 8"),
}
POV_KINI = ["JML_POV", "POV_RERATA", "POV_MEDIAN", "POV_MAKS", "POV_MIN"]

# Kertas kerja menandai sel tak berisi dengan tulisan "empty", dan pada kolom
# ukuran dengan angka 0. Keduanya berarti tidak ada nilai; disimpan sebagai
# kosong supaya tipe kolomnya tidak tercampur antara angka dan teks.
KOSONG_TEKS = {"empty", "none", "nan", "#n/a"}
KOL_UKURAN_NOL = {"UKR_MEDIA", "UKR_BPB", "UKR_STUDI", "UKR_AKP", "UKR_SK226",
                  "UKR_KORID", "UKR_CATAT"}


def bersihkan(kode, v):
    """Kembalikan None bila nilai itu sebenarnya penanda kosong."""
    if isinstance(v, str):
        s = v.strip()
        if s == "" or s.lower() in KOSONG_TEKS:
            return None
        return s
    if kode in KOL_UKURAN_NOL and isinstance(v, (int, float)) and float(v) == 0:
        return None
    return v


def baca_gpkg(path):
    con = sqlite3.connect(path)
    con.row_factory = sqlite3.Row
    tabel = con.execute("SELECT table_name FROM gpkg_contents "
                        "WHERE data_type='features'").fetchone()[0]
    tipe = {r[1]: (r[2] or "TEXT").upper()
            for r in con.execute('PRAGMA table_info("%s")' % tabel)}
    out = []
    for r in con.execute('SELECT * FROM "%s"' % tabel):
        d = dict(r)
        pt = gpkg_point(d.pop("geom", None))
        d.pop("fid", None)
        out.append({"props": {k: v for k, v in d.items() if v is not None},
                    "geom": [round(pt[0], 8), round(pt[1], 8)] if pt else None})
    con.close()
    tipe.pop("geom", None)
    tipe.pop("fid", None)
    return tipe, out


def tulis(path, obj, pretty=False):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2 if pretty else None,
                  separators=None if pretty else (",", ":"))
    return os.path.getsize(path)


def main():
    src_gpkg, src_xl = sys.argv[1], sys.argv[2]
    kamus = baca_kamus(SRC_KAMUS)
    rev = {v.strip().lower(): k for k, v in kamus.items()}

    tipe_teknis, teknis = baca_gpkg(src_gpkg)
    print("sumber teknis : %d titik, %d kolom" % (len(teknis), len(tipe_teknis)))

    wb = openpyxl.load_workbook(src_xl, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}
    xl = {}
    for r in it:
        u = r[idx[KUNCI_XL]]
        if u:
            xl[str(u)] = r
    wb.close()
    print("sumber tarif  : %d baris, %d kolom" % (len(xl), len(hdr)))

    # ---- petakan judul Excel -> kode kolom
    peta = OrderedDict()
    label = {}
    tak_dikenal = []
    for h in hdr:
        if h in ABAIKAN:
            continue
        if h in POV_POTRET:
            kode, lab = POV_POTRET[h]
        elif h in KODE_TAMBAHAN:
            kode, lab = KODE_TAMBAHAN[h]
        else:
            kode = rev.get(str(h).strip().lower())
            lab = str(h)
            if not kode:
                tak_dikenal.append(h)
                continue
        peta[h] = kode
        label[kode] = lab
    print("kolom tarif yang dibawa: %d | diabaikan (sudah ada di sumber teknis): %d"
          % (len(peta), len(ABAIKAN)))
    if tak_dikenal:
        print("  ! tidak dikenali:", tak_dikenal)

    bentrok = [k for k in peta.values() if k in tipe_teknis]
    if bentrok:
        print("  ! kode bentrok dengan sumber teknis:", bentrok)

    # ---- nilai POV terkini, dipertahankan dari data sebelumnya
    lama = {f["properties"].get("UID"): f["properties"]
            for f in json.load(open(os.path.join(DATA, "reklame.geojson"),
                                    encoding="utf-8"))["features"]}

    # ---- gabungkan
    feats = []
    tanpa_tarif = []
    for t in teknis:
        p = dict(t["props"])
        uid = p.get("UID")
        r = xl.get(uid)
        if r is None:
            tanpa_tarif.append(uid)
        else:
            for h, kode in peta.items():
                v = bersihkan(kode, r[idx[h]])
                if v is None:
                    continue
                if kode.startswith(("SEWA_NS", "PJK_RKLM", "NSR_PENUH", "PENYUSUT",
                                    "NJOP_", "UWT_KOM", "TRF_NS_HR", "BIAYA_",
                                    "TTL_LISTR", "PMLHRAAN", "ASURANSI", "PERIZINAN",
                                    "LANTAI_TN", "TTL_NS5", "PITA_CPM", "KONTAK_RB")):
                    try:
                        v = int(round(float(v)))       # rupiah tanpa pecahan
                    except (TypeError, ValueError):
                        pass
                p[kode] = v
        for c in POV_KINI:                              # POV dari hitungan terkini
            v = (lama.get(uid) or {}).get(c)
            if v is not None:
                p[c] = v
            else:
                p.pop(c, None)
        feats.append({"type": "Feature", "id": p.get("ID_TITIK"), "properties": p,
                      "geometry": {"type": "Point", "coordinates": t["geom"]}
                      if t["geom"] else None})
    if tanpa_tarif:
        print("  ! titik tanpa baris tarif: %d %s" % (len(tanpa_tarif), tanpa_tarif[:5]))

    hilang = sorted(set(xl) - {f["properties"].get("UID") for f in feats})
    if hilang:
        print("  ! baris tarif tanpa titik: %d %s" % (len(hilang), hilang[:5]))

    n1 = tulis(os.path.join(DATA, "reklame.geojson"),
               {"type": "FeatureCollection", "features": feats})

    # ---- schema.json
    for kode, lab in label.items():
        kamus.setdefault(kode, lab)
        kamus[kode] = lab
    kamus.setdefault("POV_MEDIAN", "POV median GIS (m)")
    pv = json.load(open(os.path.join(DATA, "pov.geojson"), encoding="utf-8"))["features"]

    def kolom(props_list, asli):
        urut = OrderedDict.fromkeys(k for p in props_list for k in p)
        out = []
        for c in urut:
            if c in asli:
                out.append((c, asli[c]))
                continue
            nilai = [p[c] for p in props_list if p.get(c) is not None]
            if nilai and all(isinstance(v, int) and not isinstance(v, bool) for v in nilai):
                out.append((c, "INTEGER"))
            elif nilai and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                               for v in nilai):
                out.append((c, "REAL"))
            else:
                out.append((c, "TEXT"))
        return out

    from hitung_pov import tipe_asli, SRC_GPKG_PV
    asli_rk = dict(tipe_teknis)
    asli_pv = tipe_asli(SRC_GPKG_PV)
    prk = [f["properties"] for f in feats]
    ppv = [f["properties"] for f in pv]
    skema = {
        "generated": "tools/gabung_teknis_tarif.py",
        "groups": OrderedDict([(k, l) for k, l, _ in GRUP] + [GRUP_LAIN]),
        "reklame": {"idField": "ID_TITIK",
                    "fields": bangun_skema(kolom(prk, asli_rk),
                                           [{"properties": p} for p in prk], kamus)},
        "pov": {"idField": "ID_POV", "parentField": "ID_TITIK",
                "fields": bangun_skema(kolom(ppv, asli_pv),
                                       [{"properties": p} for p in ppv], kamus)},
    }
    tulis(os.path.join(DATA, "schema.json"), skema, pretty=True)

    # ---- laporan kolom yang tidak lagi terbawa
    kol_baru = {k for p in prk for k in p}
    kol_lama = {k for p in lama.values() for k in p}
    lepas = sorted(kol_lama - kol_baru)
    print()
    print("reklame.geojson: %d fitur, %d kolom, %.2f MB"
          % (len(feats), len(skema["reklame"]["fields"]), n1 / 1e6))
    print("kolom lama yang tidak lagi terbawa (%d):" % len(lepas))
    for c in lepas:
        print("   ", c)


if __name__ == "__main__":
    main()
