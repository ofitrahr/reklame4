# -*- coding: utf-8 -*-
"""
Gabungkan nilai sewa hasil hitungan tim tarif ke data spasial.

    python tools/join_tarif.py "<berkas tarif>.xlsx"

Penggabungan memakai UID sebagai kunci. Kolom baru ditambahkan, kolom
SEWA_NS1..NS7 yang lama tidak disentuh karena berasal dari putaran
perhitungan sebelumnya.
"""
import json
import os
import sys
from collections import OrderedDict

import openpyxl

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
sys.path.insert(0, HERE)
from gpkg_to_geojson import baca_kamus, bangun_skema, GRUP, GRUP_LAIN, SRC_KAMUS
from hitung_pov import tipe_asli, SRC_GPKG_RK, SRC_GPKG_PV

# (judul kolom di Excel, nama kolom di data spasial, label untuk formulir)
PETA = [
    ("SEWA NS A", "SEWA_NSA", "Sewa NS A (Rp/th)"),
    ("SEWA NS B", "SEWA_NSB", "Sewa NS B (Rp/th)"),
    ("SEWA NS C", "SEWA_NSC", "Sewa NS C (Rp/th)"),
    ("SEWA NS D", "SEWA_NSD", "Sewa NS D (Rp/th)"),
]
KUNCI_XL = "UID GIS"


def baca_tarif(path):
    wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    it = ws.iter_rows(values_only=True)
    hdr = list(next(it))
    idx = {h: i for i, h in enumerate(hdr)}
    for xl, _gis, _label in PETA:
        if xl not in idx:
            raise SystemExit("Kolom '%s' tidak ada di berkas tarif." % xl)
    out = {}
    for r in it:
        uid = r[idx[KUNCI_XL]]
        if not uid:
            continue
        out[str(uid)] = {gis: r[idx[xl]] for xl, gis, _l in PETA}
    wb.close()
    return out


def tulis(path, obj, pretty=False):
    with open(path, "w", encoding="utf-8") as f:
        json.dump(obj, f, ensure_ascii=False, indent=2 if pretty else None,
                  separators=None if pretty else (",", ":"))
    return os.path.getsize(path)


def main():
    src = sys.argv[1]
    tarif = baca_tarif(src)
    print("berkas tarif: %d baris" % len(tarif))

    gj = json.load(open(os.path.join(DATA, "reklame.geojson"), encoding="utf-8"))
    feats = gj["features"]
    uid_gis = {f["properties"].get("UID") for f in feats}
    print("data spasial: %d titik" % len(feats))

    hanya_xl = sorted(set(tarif) - uid_gis)
    hanya_gis = sorted(uid_gis - set(tarif))
    print("UID hanya di berkas tarif : %d %s" % (len(hanya_xl), hanya_xl[:5]))
    print("UID hanya di data spasial : %d %s" % (len(hanya_gis), hanya_gis[:5]))

    terisi = OrderedDict((g, 0) for _x, g, _l in PETA)
    tanpa = []
    for f in feats:
        uid = f["properties"].get("UID")
        nilai = tarif.get(uid)
        if nilai is None:
            tanpa.append(uid)
            continue
        for _xl, gis, _l in PETA:
            v = nilai[gis]
            if v is None:
                continue
            # rupiah tidak berpecahan; dibulatkan agar tidak menyimpan angka semu
            f["properties"][gis] = int(round(float(v)))
            terisi[gis] += 1
    print()
    for g, n in terisi.items():
        print("  %-9s terisi pada %d titik" % (g, n))
    if tanpa:
        print("  titik tanpa padanan di berkas tarif: %d" % len(tanpa))

    n1 = tulis(os.path.join(DATA, "reklame.geojson"),
               {"type": "FeatureCollection", "features": feats})

    # ---- susun ulang schema.json agar kolom baru dikenali formulir
    pv = json.load(open(os.path.join(DATA, "pov.geojson"), encoding="utf-8"))["features"]
    kamus = baca_kamus(SRC_KAMUS)
    kamus.setdefault("POV_MEDIAN", "POV median GIS (m)")
    for _xl, gis, label in PETA:
        kamus[gis] = label
    asli_rk = tipe_asli(SRC_GPKG_RK)
    asli_rk.setdefault("POV_MEDIAN", "REAL")
    for _xl, gis, _l in PETA:
        asli_rk[gis] = "MEDIUMINT"
    asli_pv = tipe_asli(SRC_GPKG_PV)

    def kolom(props_list, asli):
        urut = OrderedDict.fromkeys(k for p in props_list for k in p)
        out = []
        for c in urut:
            if c in asli:
                out.append((c, asli[c]))
                continue
            nilai = [p[c] for p in props_list if p.get(c) is not None]
            if nilai and all(isinstance(v, int) and not isinstance(v, bool) for v in nilai):
                out.append((c, "INTEGER"))
            elif nilai and all(isinstance(v, (int, float)) and not isinstance(v, bool)
                               for v in nilai):
                out.append((c, "REAL"))
            else:
                out.append((c, "TEXT"))
        return out

    prk = [f["properties"] for f in feats]
    ppv = [f["properties"] for f in pv]
    skema = {
        "generated": "tools/join_tarif.py",
        "groups": OrderedDict([(k, l) for k, l, _ in GRUP] + [GRUP_LAIN]),
        "reklame": {"idField": "ID_TITIK",
                    "fields": bangun_skema(kolom(prk, asli_rk),
                                           [{"properties": p} for p in prk], kamus)},
        "pov": {"idField": "ID_POV", "parentField": "ID_TITIK",
                "fields": bangun_skema(kolom(ppv, asli_pv),
                                       [{"properties": p} for p in ppv], kamus)},
    }
    tulis(os.path.join(DATA, "schema.json"), skema, pretty=True)
    print()
    print("reklame.geojson: %d fitur, %d kolom, %.2f MB"
          % (len(feats), len(skema["reklame"]["fields"]), n1 / 1e6))


if __name__ == "__main__":
    main()
