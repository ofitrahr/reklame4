# -*- coding: utf-8 -*-
"""
Ekspor untuk pemberi kerja: satu Shapefile berisi kolom teknis dan keempat
kolom sewa, serta empat berkas KMZ yang masing-masing memuat kolom teknis
dan satu susunan sewa.

    python tools/ekspor_shp_kmz.py

Shapefile dibuat lewat ogr2ogr bawaan QGIS agar berkas pendampingnya
(.prj, .cpg, .dbf) sesuai kaidah. KMZ disusun langsung supaya titiknya
berwarna menurut besaran sewa dan balon keterangannya rapi dibaca.
"""
import json
import os
import re
import shutil
import subprocess
import sys
import zipfile
from xml.sax.saxutils import escape

HERE = os.path.dirname(os.path.abspath(__file__))
ROOT = os.path.dirname(HERE)
DATA = os.path.join(ROOT, "data")
KELUAR = os.path.join(ROOT, "keluaran")

# (kolom di data, judul untuk pemberi kerja)
TEKNIS = [
    ("UID", "UID"),
    ("ID_TITIK", "ID Titik"),
    ("KEWENANGAN", "Kewenangan"),
    ("JENIS", "Jenis"),
    ("TIPE", "Tipe"),
    ("TIPE_MEDIA", "Tipe Media"),
    ("UKR_MEDIA", "Dimensi Ukuran"),
    ("MNMPL_STR", "Bentuk (menempel struktur)"),
    ("JML_MUKA", "Jumlah Muka"),
    ("NAMA_JALAN", "Nama Jalan"),
    ("KELAS_JLN", "Kelas Jalan"),
    ("KELURAHAN", "Kelurahan"),
    ("KECAMATAN", "Kecamatan"),
    ("STATUS_WPP", "Sub Wilayah SK 321"),
    ("POV_MEDIAN", "POV Median (m)"),
    ("ROW_JALAN", "ROW Jalan (m)"),
]
SEWA = [("SEWA_NSA", "Sewa NS A (Rp/th)"),
        ("SEWA_NSB", "Sewa NS B (Rp/th)"),
        ("SEWA_NSC", "Sewa NS C (Rp/th)"),
        ("SEWA_NSD", "Sewa NS D (Rp/th)")]

OGR2OGR = r"C:\Program Files\QGIS 3.34.12\bin\ogr2ogr.exe"

# Warna kelas sewa, dari paling rendah ke paling tinggi (KML: aabbggrr).
WARNA = ["ffb2ffff", "ff5cccfe", "ff3c8dfd", "ff203bf0", "ff2600bd"]
WARNA_NOL = "ff9e9e9e"
IKON = "http://maps.google.com/mapfiles/kml/shapes/placemark_circle.png"


def muat():
    gj = json.load(open(os.path.join(DATA, "reklame.geojson"), encoding="utf-8"))
    return gj["features"]


def rupiah(v):
    return "-" if v is None else "Rp " + format(int(v), ",d").replace(",", ".")


# ------------------------------------------------------------------ shapefile
def buat_shp(feats, awalan):
    kolom = [k for k, _j in TEKNIS] + [k for k, _j in SEWA]
    panjang = [k for k in kolom if len(k) > 10]
    if panjang:
        raise SystemExit("Nama kolom melebihi 10 huruf: %s" % panjang)

    tmp = os.path.join(KELUAR, "_shp_sumber.geojson")
    out_dir = os.path.join(KELUAR, "%s_SHP" % awalan)
    shutil.rmtree(out_dir, ignore_errors=True)
    os.makedirs(out_dir, exist_ok=True)

    ringkas = {"type": "FeatureCollection", "features": [
        {"type": "Feature",
         "properties": {k: f["properties"].get(k) for k in kolom},
         "geometry": f["geometry"]} for f in feats]}
    with open(tmp, "w", encoding="utf-8") as fh:
        json.dump(ringkas, fh, ensure_ascii=False)

    shp = os.path.join(out_dir, "%s_CLIENT.shp" % awalan)
    cmd = [OGR2OGR, "-f", "ESRI Shapefile", shp, tmp,
           "-nln", "TITIK_REKLAME", "-a_srs", "EPSG:4326",
           "-lco", "ENCODING=UTF-8"]
    # Instalasi PostgreSQL/PostGIS di komputer ini memasang PROJ_LIB ke pustaka
    # PROJ versi lama, yang membuat ogr2ogr QGIS menolak definisi EPSG. Arahkan
    # ke pustaka milik QGIS khusus untuk pemanggilan ini.
    env = dict(os.environ)
    proj_qgis = os.path.join(os.path.dirname(os.path.dirname(OGR2OGR)), "share", "proj")
    if os.path.exists(os.path.join(proj_qgis, "proj.db")):
        env["PROJ_LIB"] = proj_qgis
        env["PROJ_DATA"] = proj_qgis
    r = subprocess.run(cmd, capture_output=True, text=True, env=env)
    if r.returncode != 0:
        raise SystemExit("ogr2ogr gagal:\n" + r.stdout + r.stderr)
    os.remove(tmp)

    berkas = sorted(os.listdir(out_dir))
    print("Shapefile -> %s" % out_dir)
    for b in berkas:
        print("   %-34s %8.1f KB" % (b, os.path.getsize(os.path.join(out_dir, b)) / 1024))
    return out_dir, kolom


# ------------------------------------------------------------------ kmz
def kelas(nilai, batas):
    for i, b in enumerate(batas):
        if nilai <= b:
            return i
    return len(batas)


def buat_kmz(feats, kunci, judul, awalan):
    isi = [f["properties"].get(kunci) for f in feats]
    positif = sorted(v for v in isi if v)
    # lima kelas berdasarkan kuantil dari titik yang bernilai
    batas = [positif[int(len(positif) * q)] for q in (0.2, 0.4, 0.6, 0.8)] if positif else [0]

    baris = []
    baris.append('<?xml version="1.0" encoding="UTF-8"?>')
    baris.append('<kml xmlns="http://www.opengis.net/kml/2.2">')
    baris.append("<Document>")
    baris.append("<name>%s</name>" % escape(judul))
    ket = ("Titik reklame Kota Batam. Warna titik menurut %s: "
           "abu-abu = tidak ditarifkan; kuning muda sampai merah tua = "
           "nilai terendah sampai tertinggi." % escape(judul))
    baris.append("<description><![CDATA[%s]]></description>" % ket)

    balon = ("<BalloonStyle><text><![CDATA["
             "<h3>$[name]</h3><table cellpadding='3' style='font-family:Arial;font-size:12px'>"
             + "".join("<tr><td><b>%s</b></td><td>$[%s]</td></tr>" % (escape(j), k)
                       for k, j in TEKNIS)
             + "<tr><td><b>%s</b></td><td>$[%s]</td></tr>" % (escape(judul), kunci)
             + "</table>]]></text></BalloonStyle>")
    for i, w in enumerate(WARNA + [WARNA_NOL]):
        baris.append(
            '<Style id="k%d"><IconStyle><color>%s</color><scale>0.9</scale>'
            '<Icon><href>%s</href></Icon></IconStyle>'
            '<LabelStyle><scale>0</scale></LabelStyle>%s</Style>'
            % (i, w, IKON, balon))

    per_kec = {}
    for f in feats:
        per_kec.setdefault(f["properties"].get("KECAMATAN") or "TANPA KECAMATAN",
                           []).append(f)

    for kec in sorted(per_kec):
        baris.append("<Folder><name>%s (%d)</name>" % (escape(kec), len(per_kec[kec])))
        for f in per_kec[kec]:
            p = f["properties"]
            g = f["geometry"]["coordinates"] if f["geometry"] else None
            if not g:
                continue
            v = p.get(kunci)
            gaya = len(WARNA) if not v else kelas(v, batas)
            baris.append("<Placemark>")
            baris.append("<name>%s</name>" % escape(str(p.get("ID_TITIK") or "")))
            baris.append("<styleUrl>#k%d</styleUrl>" % gaya)
            baris.append("<ExtendedData>")
            for k, j in TEKNIS:
                nilai = p.get(k)
                teks = "-" if nilai is None else str(nilai)
                baris.append('<Data name="%s"><displayName>%s</displayName>'
                             "<value>%s</value></Data>"
                             % (k, escape(j), escape(teks)))
            baris.append('<Data name="%s"><displayName>%s</displayName>'
                         "<value>%s</value></Data>"
                         % (kunci, escape(judul), escape(rupiah(v))))
            baris.append("</ExtendedData>")
            baris.append("<Point><coordinates>%.8f,%.8f,0</coordinates></Point>"
                         % (g[0], g[1]))
            baris.append("</Placemark>")
        baris.append("</Folder>")
    baris.append("</Document></kml>")

    nama = "%s_%s.kmz" % (awalan, kunci.replace("SEWA_", ""))
    jalur = os.path.join(KELUAR, "KMZ", nama)
    os.makedirs(os.path.dirname(jalur), exist_ok=True)
    with zipfile.ZipFile(jalur, "w", zipfile.ZIP_DEFLATED) as z:
        z.writestr("doc.kml", "\n".join(baris))
    print("   %-30s %7.2f MB | kelas: %s"
          % (nama, os.path.getsize(jalur) / 1e6,
             " / ".join(rupiah(b) for b in batas)))
    return jalur


# ------------------------------------------------------------------ excel
def buat_xlsx(feats, nama):
    """Lembar kerja dengan susunan kolom yang sama seperti Shapefile."""
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.utils import get_column_letter
    from datetime import datetime

    kolom = TEKNIS + SEWA
    lebar = {"ID_TITIK": 24, "NAMA_JALAN": 38, "STATUS_WPP": 26, "KELURAHAN": 20,
             "TIPE": 22, "KEWENANGAN": 17, "MNMPL_STR": 24}
    wb = Workbook()
    ws = wb.active
    ws.title = "Titik Reklame"
    kf = Font(name="Arial", size=10, bold=True, color="FFFFFF")
    ki = PatternFill("solid", fgColor="1F4E79")
    gs = Side(style="thin", color="BFBFBF")
    tepi = Border(left=gs, right=gs, top=gs, bottom=gs)
    isi_font = Font(name="Arial", size=10)

    for i, (k, j) in enumerate(kolom, start=1):
        c = ws.cell(row=1, column=i, value=j)
        c.font, c.fill, c.border = kf, ki, tepi
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = lebar.get(k, 15)
    ws.row_dimensions[1].height = 30

    for r, f in enumerate(feats, start=2):
        p = f["properties"]
        for i, (k, _j) in enumerate(kolom, start=1):
            c = ws.cell(row=r, column=i, value=p.get(k))
            c.font, c.border = isi_font, tepi
            if k in ("POV_MEDIAN", "ROW_JALAN"):
                c.number_format = "0.0"
                c.alignment = Alignment(horizontal="right")
            elif k.startswith("SEWA_") or k == "JML_MUKA":
                c.number_format = "#,##0"
                c.alignment = Alignment(horizontal="right")
    ws.freeze_panes = "C2"
    ws.auto_filter.ref = "A1:%s%d" % (get_column_letter(len(kolom)), len(feats) + 1)

    k = wb.create_sheet("Keterangan")
    k.column_dimensions["A"].width = 26
    k.column_dimensions["B"].width = 96
    k.cell(row=1, column=1, value="Keterangan Berkas").font = Font(name="Arial", size=12, bold=True)
    isi = [("Berkas", "Titik reklame Kota Batam - susunan kolom permintaan pemberi kerja"),
           ("Jumlah titik", len(feats)),
           ("Dibuat", datetime.now().strftime("%d %B %Y, %H:%M")),
           ("Sistem koordinat", "WGS 84 (EPSG:4326) - lihat berkas SHP/GPKG untuk geometrinya"),
           ("", ""),
           ("Bentuk", "Diambil dari kolom MNMPL_STR, yaitu apakah reklame menempel pada "
                      "struktur bangunan atau berdiri sendiri."),
           ("Dimensi Ukuran", "Ukuran bidang yang dipakai dalam perhitungan tarif NS A-D."),
           ("POV Median", "Nilai tengah jarak titik POV ke reklame, dalam meter. Kosong "
                          "berarti titik itu belum memiliki titik POV."),
           ("ROW Jalan", "Lebar ruang milik jalan hasil pengukuran tim GIS. Hanya diukur "
                         "untuk titik koridor sisi jalan."),
           ("Sewa NS A - NS D", "Rupiah per tahun. Nilai 0 berarti titik tidak ditarifkan "
                                "pada susunan tersebut; sel kosong berarti tidak dihitung.")]
    for i, (a, b) in enumerate(isi, start=3):
        ca = k.cell(row=i, column=1, value=a)
        ca.font = Font(name="Arial", size=10, bold=True)
        ca.alignment = Alignment(vertical="top")
        cb = k.cell(row=i, column=2, value=b)
        cb.font = Font(name="Arial", size=10)
        cb.alignment = Alignment(vertical="top", wrap_text=True)

    jalur = os.path.join(KELUAR, nama)
    wb.save(jalur)
    print("Excel     -> %s  (%d baris x %d kolom, %.2f MB)"
          % (nama, len(feats), len(kolom), os.path.getsize(jalur) / 1e6))
    return jalur


def main():
    awalan = sys.argv[1] if len(sys.argv) > 1 else "Reklame_Batam"
    os.makedirs(KELUAR, exist_ok=True)
    feats = muat()
    print("data: %d titik" % len(feats))
    print()
    buat_shp(feats, awalan)
    print()
    buat_xlsx(feats, "%s_CLIENT.xlsx" % awalan)
    print()
    print("KMZ -> %s" % os.path.join(KELUAR, "KMZ"))
    for kunci, judul in SEWA:
        buat_kmz(feats, kunci, judul, awalan)


if __name__ == "__main__":
    main()
